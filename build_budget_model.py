from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

F="Arial"
BLUE=Font(name=F,color="0000FF"); BLK=Font(name=F); GRN=Font(name=F,color="008000")
B=Font(name=F,bold=True); H=Font(name=F,bold=True,size=14); W=Font(name=F,bold=True,color="FFFFFF")
YEL=PatternFill("solid",fgColor="FFFF00"); HDR=PatternFill("solid",fgColor="1F3864"); SEC=PatternFill("solid",fgColor="D9E1F2")
thin=Side(style="thin",color="999999"); TOP=Border(top=thin)
CUR='$#,##0;($#,##0);"-"'; CUR2='$#,##0.00;($#,##0.00);"-"'; PCT='0.0%;(0.0%);"-"'; NUM='#,##0;(#,##0);"-"'; NUM1='#,##0.0;(#,##0.0);"-"'
wb=Workbook()

def section(ws,r,text,ncols=6):
    for c in range(1,ncols+1): ws.cell(r,c).fill=SEC
    ws.cell(r,1,text).font=B
def hdr(ws,r,vals):
    for i,v in enumerate(vals,1):
        c=ws.cell(r,i,v); c.font=W; c.fill=HDR; c.alignment=Alignment(horizontal="center",wrap_text=True,vertical="center")
def inp(c,v,fmt=None,key=True):
    c.value=v; c.font=BLUE
    if key: c.fill=YEL
    if fmt: c.number_format=fmt

# ---------------- Inputs ----------------
ws=wb.active; ws.title="Inputs"
ws["A1"]="2027 Marketing Budget Model: Inputs"; ws["A1"].font=H
ws["A2"]="Yellow cells with blue text are inputs. Change them and every other tab recalculates."; ws["A2"].font=Font(name=F,italic=True)
for col,w in zip("ABCDEFG",[44,16,14,16,16,16,60]): ws.column_dimensions[col].width=w
R={}
r=4; section(ws,r,"Property profile",7); r+=1
rows=[("name","Property name","Harrison Landing",None,"","From Property List (Haven Homes)"),
("port","Portfolio","Haven Homes",None,"",""),
("loc","Location","Simpsonville, SC (Greenville Co.)",None,"","From Property List"),
("type","Property type / ownership","Townhome / Owned-Managed",None,"",""),
("units","Total units",166,NUM,"units","From Property List"),
("rent","Average monthly rent",2100,CUR,"$/mo","PLACEHOLDER: replace with rent roll average"),
("occ","Current occupancy",0.95,PCT,"%","PLACEHOLDER: current physical occupancy from PMS"),
("tocc","Target occupancy (end of 2027)",0.96,PCT,"%","PLACEHOLDER: owner/asset mgmt target"),
("dv","Average days vacant per turn",30,NUM,"days","PLACEHOLDER: move-out to move-in, from PMS turn report"),]
def put(rows,r):
    for k,lab,v,fmt,unit,note in rows:
        ws.cell(r,1,lab).font=BLK
        inp(ws.cell(r,2),v,fmt,key=not isinstance(v,str) or k in("name",))
        if isinstance(v,str): ws.cell(r,2).font=BLUE
        ws.cell(r,3,unit).font=BLK; ws.cell(r,7,note).font=Font(name=F,italic=True,color="595959")
        R[k]=f"Inputs!$B${r}"; r+=1
    return r
r=put(rows,r); r+=1
section(ws,r,"Lease expirations & retention",7); r+=1
r=put([("exp","% of occupied leases expiring in 2027",0.95,PCT,"%","PLACEHOLDER: count 2027 expirations in lease expiration report ÷ occupied units"),
("ren","Renewal rate (of expiring leases)",0.55,PCT,"%","PLACEHOLDER: 2026 YTD renewals ÷ expirations"),
("et","Early terminations / skips (% of units per year)",0.03,PCT,"%","PLACEHOLDER: breaks, skips, evictions, transfers out"),
("ri","Renewal incentive per renewal",100,CUR,"$","PLACEHOLDER: gift card / upgrade offered to renew; enter 0 if none"),],r); r+=1
section(ws,r,"Leasing funnel (conversion rates)",7); r+=1
r=put([("c1","Lead → tour",0.25,PCT,"%","PLACEHOLDER: pull from CRM funnel report"),
("c2","Tour → application",0.35,PCT,"%","PLACEHOLDER: includes self-guided tours"),
("c3","Application → approval",0.70,PCT,"%","PLACEHOLDER: screening approval rate"),
("c4","Approval → signed lease",0.90,PCT,"%","PLACEHOLDER: approved applicants who sign"),],r)
ws.cell(r,1,"Lead-to-lease rate (calculated)").font=B
ws.cell(r,2,f"={R['c1']}*{R['c2']}*{R['c3']}*{R['c4']}").number_format=PCT; ws.cell(r,2).font=B
ws.cell(r,7,"Industry range is often ~3–8%; your real number matters most").font=Font(name=F,italic=True,color="595959")
R["ltl"]=f"Inputs!$B${r}"; r+=2
section(ws,r,"Budget guardrails",7); r+=1
r=put([("cont","Contingency / reserve",0.10,PCT,"%","Cushion for slow months, price hikes, surprise vacancies"),
("cap","Target max cost per lease (% of one month's rent)",0.50,PCT,"%","Common rule of thumb: keep acquisition cost under ~½ month's rent"),],r); r+=1

# Seasonality
section(ws,r,"Lease expiration seasonality (% of 2027 expirations by month)",7); r+=1
hdr(ws,r,["Month","% of expirations","","","","","Notes"]); r+=1
months=["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
season=[.05,.05,.07,.09,.11,.13,.13,.11,.08,.07,.06,.05]
s0=r
for m,v in zip(months,season):
    ws.cell(r,1,m).font=BLK; inp(ws.cell(r,2),v,PCT); r+=1
ws.cell(s0,7,"PLACEHOLDER curve (summer-heavy). Replace with actual 2027 expiration counts by month.").font=Font(name=F,italic=True,color="595959")
ws.cell(r,1,"Total (must = 100%)").font=B
ws.cell(r,2,f"=SUM(B{s0}:B{r-1})").number_format=PCT; ws.cell(r,2).font=B
ws.cell(r,3,f'=IF(ROUND(B{r},4)=1,"OK","CHECK")').font=B
R["seas0"]=s0; R["seasChk"]=f"Inputs!$C${r}"; r+=2

# Channels
section(ws,r,"Lead channels",7); r+=1
hdr(ws,r,["Channel","% of leads","Cost per lead ($)","Fixed monthly fee ($)","Cost per signed lease ($)","2026 actual spend ($)","Notes"]); r+=1
ch=[("Zillow Rentals",0.30,25,0,0,"PLACEHOLDER: check your Zillow contract (per-lead vs. flat)"),
("Apartments.com / CoStar network",0.20,0,1200,0,"PLACEHOLDER: flat package; confirm 2027 renewal price"),
("Google paid search",0.15,35,0,0,"PLACEHOLDER: brand + 'townhomes for rent Simpsonville'"),
("Paid social (Meta)",0.10,20,0,0,"PLACEHOLDER"),
("Website / SEO (organic)",0.10,0,0,0,"Website cost lives in fixed costs below"),
("Resident referrals",0.05,0,0,300,"PLACEHOLDER: referral bonus paid per signed lease"),
("Signage / drive-by / walk-in",0.10,0,0,0,"Signage cost lives in fixed costs below"),]
c0=r
for name,sh,cpl,fx,cpls,note in ch:
    ws.cell(r,1,name).font=BLUE; ws.cell(r,1).fill=YEL
    inp(ws.cell(r,2),sh,PCT); inp(ws.cell(r,3),cpl,CUR); inp(ws.cell(r,4),fx,CUR); inp(ws.cell(r,5),cpls,CUR)
    inp(ws.cell(r,6),None,CUR)
    ws.cell(r,7,note).font=Font(name=F,italic=True,color="595959"); r+=1
c1=r-1
ws.cell(r,1,"Total (must = 100%)").font=B
ws.cell(r,2,f"=SUM(B{c0}:B{c1})").number_format=PCT; ws.cell(r,2).font=B
ws.cell(r,3,f'=IF(ROUND(B{r},4)=1,"OK","CHECK")').font=B
R["chChk"]=f"Inputs!$C${r}"
ws.cell(r+1,1,"Leave 2026 actual blank if unknown; the Summary tab will skip the variance.").font=Font(name=F,italic=True,color="595959")
r+=3

# Fixed costs
section(ws,r,"Fixed & program costs (property-level)",7); r+=1
hdr(ws,r,["Item","Monthly ($)","One-time / annual ($)","","","","Notes"]); r+=1
fx=[("CRM / lead management (allocated)",250,0,"PLACEHOLDER: e.g. ~$1.50/unit/mo"),
("AI leasing assistant / chat",166,0,"PLACEHOLDER: e.g. ~$1/unit/mo"),
("Reputation management / review software",150,0,"PLACEHOLDER"),
("Website hosting & maintenance",100,0,"PLACEHOLDER"),
("Photography / video / virtual tour refresh",0,1500,"PLACEHOLDER: one-time 2027 refresh"),
("Signage & banners",0,1000,"PLACEHOLDER"),
("Print collateral",0,300,"PLACEHOLDER"),
("Resident events & retention",0,3000,"PLACEHOLDER: ~4 events/yr; retention is cheaper than replacement"),
("Freelance / agency support",0,0,"Optional: outsourced help for a one-person team"),]
f0=r
for name,mo,ot,note in fx:
    ws.cell(r,1,name).font=BLUE; ws.cell(r,1).fill=YEL
    inp(ws.cell(r,2),mo,CUR); inp(ws.cell(r,3),ot,CUR)
    ws.cell(r,7,note).font=Font(name=F,italic=True,color="595959"); r+=1
f1=r-1
ws.cell(r,7,"One-time / annual amounts are spread evenly across 12 months in the Monthly Plan.").font=Font(name=F,italic=True,color="595959")
ws.freeze_panes="A4"

# ---------------- Monthly Plan ----------------
mp=wb.create_sheet("Monthly Plan")
mp["A1"]="=\"2027 Monthly Lease Demand & Spend: \"&"+R["name"]; mp["A1"].font=H
mp["A2"]="All cells are formulas driven by the Inputs tab. Assumes leads are generated in the same month as the lease need."; mp["A2"].font=Font(name=F,italic=True)
mp.column_dimensions["A"].width=42
for i in range(2,15): mp.column_dimensions[L(i)].width=11
hdr(mp,4,["Line item"]+months+["2027 Total"])
MC=lambda i:L(i+2)   # month i (0-based) column
row=5; P={}
def line(key,label,ftpl,fmt=NUM1,bold=False,total=True,font=BLK):
    global row
    mp.cell(row,1,label).font=B if bold else BLK
    for i in range(12):
        c=mp.cell(row,i+2,ftpl(i)); c.number_format=fmt; c.font=Font(name=F,bold=bold,color=font.color)
    if total:
        c=mp.cell(row,14,f"=SUM(B{row}:M{row})"); c.number_format=fmt; c.font=B
    P[key]=row; row+=1
def sec(t):
    global row
    for c in range(1,15): mp.cell(row,c).fill=SEC
    mp.cell(row,1,t).font=B; row+=1
sec("Lease demand")
expiring=f"({R['units']}*{R['occ']}*{R['exp']})"
line("sea","Expiration seasonality",lambda i:f"=Inputs!$B${R['seas0']+i}",PCT,font=GRN)
line("exp","Leases expiring",lambda i:f"={expiring}*{MC(i)}{P['sea']}")
line("ren","Renewals",lambda i:f"={MC(i)}{P['exp']}*{R['ren']}")
line("mo","Move-outs at expiration",lambda i:f"={MC(i)}{P['exp']}-{MC(i)}{P['ren']}")
line("et","Early terminations / skips",lambda i:f"={R['units']}*{R['et']}/12")
line("gap","Leases to reach target occupancy",lambda i:f"=MAX(0,{R['units']}*({R['tocc']}-{R['occ']}))/12")
line("lz","New leases needed",lambda i:f"={MC(i)}{P['mo']}+{MC(i)}{P['et']}+{MC(i)}{P['gap']}",bold=True)
sec("Funnel targets")
line("leads","Leads needed",lambda i:f"=IFERROR({MC(i)}{P['lz']}/{R['ltl']},0)",NUM,bold=True)
line("tours","Tours needed",lambda i:f"={MC(i)}{P['leads']}*{R['c1']}",NUM)
line("apps","Applications needed",lambda i:f"={MC(i)}{P['tours']}*{R['c2']}",NUM)
sec("Spend by lead channel ($)")
chrows=[]
for k in range(c0,c1+1):
    line(f"ch{k}",f"=Inputs!$A${k}",lambda i,k=k:(f"=Inputs!$D${k}+{MC(i)}{P['leads']}*Inputs!$B${k}*Inputs!$C${k}"
        f"+{MC(i)}{P['lz']}*Inputs!$B${k}*Inputs!$E${k}"),CUR)
    mp.cell(P[f"ch{k}"],1).font=GRN; chrows.append(P[f"ch{k}"])
line("chsub","Channel subtotal",lambda i:f"=SUM({MC(i)}{chrows[0]}:{MC(i)}{chrows[-1]})",CUR,bold=True)
sec("Fixed, program & retention costs ($)")
fxrows=[]
for k in range(f0,f1+1):
    line(f"fx{k}",f"=Inputs!$A${k}",lambda i,k=k:f"=Inputs!$B${k}+Inputs!$C${k}/12",CUR)
    mp.cell(P[f"fx{k}"],1).font=GRN; fxrows.append(P[f"fx{k}"])
line("ri","Renewal incentives",lambda i:f"={MC(i)}{P['ren']}*{R['ri']}",CUR)
line("fxsub","Fixed & retention subtotal",lambda i:f"=SUM({MC(i)}{fxrows[0]}:{MC(i)}{P['ri']})",CUR,bold=True)
sec("Total")
line("pre","Budget before contingency",lambda i:f"={MC(i)}{P['chsub']}+{MC(i)}{P['fxsub']}",CUR)
line("cont","Contingency",lambda i:f"={MC(i)}{P['pre']}*{R['cont']}",CUR)
line("tot","TOTAL MARKETING BUDGET",lambda i:f"={MC(i)}{P['pre']}+{MC(i)}{P['cont']}",CUR,bold=True)
for c in range(1,15): mp.cell(P["tot"],c).border=TOP
mp.freeze_panes="B5"

# ---------------- Budget Summary ----------------
sm=wb.create_sheet("Budget Summary",0)
sm.column_dimensions["A"].width=46; sm.column_dimensions["B"].width=18; sm.column_dimensions["C"].width=18; sm.column_dimensions["D"].width=18; sm.column_dimensions["E"].width=50
sm["A1"]="=\"2027 Marketing Budget: \"&"+R["name"]; sm["A1"].font=H
sm["A2"]="="+R["port"]+"&\" · \"&"+R["loc"]+"&\" · \"&TEXT("+R["units"]+",\"0\")&\" units\""; sm["A2"].font=Font(name=F,italic=True)
T=lambda k:f"='Monthly Plan'!$N${P[k]}"
r=4; section(sm,r,"Headline",5); r+=1
S={}
def srow(k,label,f,fmt,note="",bold=False):
    global r
    sm.cell(r,1,label).font=B if bold else BLK
    c=sm.cell(r,2,f); c.number_format=fmt; c.font=Font(name=F,bold=bold,color="008000" if "Monthly Plan" in f or "Inputs!" in f else "000000")
    sm.cell(r,5,note).font=Font(name=F,italic=True,color="595959"); S[k]=f"B{r}"; r+=1
srow("tot","Total 2027 marketing budget",T("tot"),CUR,bold=True)
srow("pu","Budget per unit per year",f"=IFERROR(B5/{R['units']},0)",CUR)
srow("pum","Budget per unit per month",f"=IFERROR(B5/{R['units']}/12,0)",CUR2)
srow("gpr","% of gross potential rent",f"=IFERROR(B5/({R['units']}*{R['rent']}*12),0)",PCT,"Many operators land ~1–3% of GPR")
r+=1; section(sm,r,"Lease demand",5); r+=1
srow("exp","Leases expiring in 2027",T("exp"),NUM)
srow("ren","Renewals",T("ren"),NUM)
srow("turns","Move-outs (expirations + early terminations)",f"='Monthly Plan'!$N${P['mo']}+'Monthly Plan'!$N${P['et']}",NUM)
srow("tr","Implied annual turnover rate",f"=IFERROR({S['turns']}/{R['units']},0)",PCT,"Move-outs ÷ total units")
srow("lz","New leases needed",T("lz"),NUM,bold=True)
srow("ltl","Lead-to-lease rate",f"={R['ltl']}",PCT)
srow("leads","Leads needed",T("leads"),NUM,bold=True)
srow("tours","Tours needed",T("tours"),NUM)
r+=1; section(sm,r,"Efficiency checks",5); r+=1
srow("cpl","Blended cost per lead",f"=IFERROR('Monthly Plan'!$N${P['chsub']}/{S['leads']},0)",CUR2,"Channel spend ÷ leads")
srow("cpls","Cost per new lease (all-in)",f"=IFERROR({S['tot']}/{S['lz']},0)",CUR,"Total budget ÷ new leases")
srow("cplr","Cost per lease as % of one month's rent",f"=IFERROR({S['cpls']}/{R['rent']},0)",PCT)
srow("capchk","Within cost-per-lease target?",f"=IF({S['cplr']}<={R['cap']},\"YES\",\"OVER TARGET\")","General",bold=True)
r+=1; section(sm,r,"Why it pays: cost of vacancy",5); r+=1
srow("vl","Projected 2027 vacancy loss from turns",f"={S['turns']}*{R['dv']}*{R['rent']}/30",CUR,"Move-outs × days vacant × daily rent")
srow("v1","Value of cutting 1 vacant day per turn",f"={S['turns']}*{R['rent']}/30",CUR,"Use this to justify spend that speeds up leasing")
srow("rv","Value of +5 pts renewal rate",f"={S['exp']}*0.05*({R['dv']}*{R['rent']}/30+IFERROR({S['cpls']},0))",CUR,"Avoided vacancy + avoided acquisition cost")
r+=1; section(sm,r,"Budget by line item",5); r+=1
hdr(sm,r,["Line item","2027 Budget ($)","2026 Actual ($)","Variance ($)","Share of budget"]); r+=1
b0=r
for k in range(c0,c1+1):
    pr=P[f"ch{k}"]
    sm.cell(r,1,f"=Inputs!$A${k}").font=GRN
    sm.cell(r,2,f"='Monthly Plan'!$N${pr}").number_format=CUR; sm.cell(r,2).font=GRN
    sm.cell(r,3,f'=IF(Inputs!$F${k}="","",Inputs!$F${k})').number_format=CUR; sm.cell(r,3).font=GRN
    sm.cell(r,4,f'=IF(C{r}="","",B{r}-C{r})').number_format=CUR
    sm.cell(r,5,f"=IFERROR(B{r}/{S['tot']},0)").number_format=PCT; r+=1
for label,k in [("Fixed, program & retention costs","fxsub"),("Contingency","cont")]:
    sm.cell(r,1,label).font=BLK
    sm.cell(r,2,T(k)).number_format=CUR; sm.cell(r,2).font=GRN
    sm.cell(r,5,f"=IFERROR(B{r}/{S['tot']},0)").number_format=PCT; r+=1
sm.cell(r,1,"Total").font=B
sm.cell(r,2,f"=SUM(B{b0}:B{r-1})").number_format=CUR; sm.cell(r,2).font=B
sm.cell(r,5,f"=SUM(E{b0}:E{r-1})").number_format=PCT; sm.cell(r,5).font=B
for c in range(1,6): sm.cell(r,c).border=TOP
r+=2; section(sm,r,"Input checks",5); r+=1
sm.cell(r,1,"Seasonality sums to 100%").font=BLK; sm.cell(r,2,"="+R["seasChk"]).font=GRN; r+=1
sm.cell(r,1,"Channel lead mix sums to 100%").font=BLK; sm.cell(r,2,"="+R["chChk"]).font=GRN; r+=1

# ---------------- Read Me ----------------
rm=wb.create_sheet("Read Me")
rm.column_dimensions["A"].width=110
lines=[("How this model works",H),
("Lease demand drives the budget: Units × occupancy × % expiring → renewals vs. move-outs → new leases needed → ÷ lead-to-lease rate → leads needed → × cost per lead by channel → + fixed costs → + contingency.",BLK),
("",BLK),("Tabs",B),
("Budget Summary: headline budget, $/unit, cost per lease, vacancy-cost justification, line items vs. 2026.",BLK),
("Inputs: the ONLY tab you edit. Property facts, retention, funnel, seasonality, channels, fixed costs.",BLK),
("Monthly Plan: month-by-month leases, leads and spend, so you can see the summer peak and pace spend.",BLK),
("",BLK),("Color legend",B),
("Blue text on yellow = input you can change · Black = formula · Green = pulled from another tab.",BLK),
("",BLK),("Status of the numbers",B),
("Units, name, location and type come from the Haven Homes property list. Everything marked PLACEHOLDER is an industry-style starting value; replace with Harrison Landing actuals.",BLK),
("",BLK),("Data to pull for Harrison Landing (replace placeholders)",B),
("1. Rent roll: average rent and current occupancy (PMS).",BLK),
("2. Lease expiration report: count of 2027 expirations by month (sets % expiring and seasonality).",BLK),
("3. 2025–2026 renewal rate and early terminations / skips.",BLK),
("4. CRM funnel report: leads, tours, applications, approvals, leases, ideally by lead source.",BLK),
("5. Invoices / contracts: Zillow, Apartments.com, Google, Meta, software; note 2027 renewal pricing.",BLK),
("6. Turn report: average days vacant (move-out to move-in).",BLK),
("",BLK),("Rolling out to the portfolio",B),
("Copy this workbook once per property (or duplicate the Inputs tab) and change the Inputs. Prosper 207 is in lease-up and needs a separate lease-up model (absorption pace, not turnover).",BLK),]
for i,(t,f) in enumerate(lines,1):
    c=rm.cell(i,1,t); c.font=f; c.alignment=Alignment(wrap_text=True)

wb.calculation.fullCalcOnLoad=True
wb.save("Harrison_Landing_2027_Marketing_Budget.xlsx")

